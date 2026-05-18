"""k-NN graph construction, Leiden clustering, UMAP embedding, and main CLI.

Follows the same structural pattern as the locomotion LocomotionAnalysisPipeline.
"""

from __future__ import annotations

import argparse
import csv
import json
import logging
import sys
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import numpy as np

from .. import paths as pathutil
from ..logging_setup import configure_logging
from ..tqdm_optional import try_tqdm
from .config import ClusteringConfig
from .dataset import BehaviorDataset, WindowMetadata
from .feature_extraction import extract_all_features, _fmt_duration, _fmt_size
from .sequence_features import SEQUENCE_FEATURE_NAMES

logger = logging.getLogger(__name__)

# Optional imports with graceful fallback
try:
    import igraph as ig
    import leidenalg
    _HAS_LEIDEN = True
except ImportError:
    _HAS_LEIDEN = False

try:
    import umap
    _HAS_UMAP = True
except ImportError:
    _HAS_UMAP = False


# ---------------------------------------------------------------------------
# Result container
# ---------------------------------------------------------------------------

@dataclass
class ClusteringResult:
    features_normalized: np.ndarray     # (N, 33)
    cluster_labels: np.ndarray          # (N,) int
    n_clusters: int
    embedding_2d: np.ndarray            # (N, 2)
    metadata: list[WindowMetadata]
    feature_names: list[str]
    method: str                         # "leiden" or "spectral"
    embedding_method: str               # "umap" or "tsne"


# ---------------------------------------------------------------------------
# Normalization
# ---------------------------------------------------------------------------

def _normalize_features(
    features: np.ndarray,
    method: str = "zscore",
) -> tuple[np.ndarray, dict[str, Any]]:
    """Normalize feature matrix, handling NaN columns gracefully.

    Returns (normalized, stats_dict).
    """
    X = features.copy()
    stats: dict[str, Any] = {"method": method}

    # Replace any remaining NaN with column median
    n_imputed = 0
    for col in range(X.shape[1]):
        mask = np.isfinite(X[:, col])
        if not np.all(mask):
            nan_count = int(np.sum(~mask))
            median_val = float(np.nanmedian(X[:, col])) if np.any(mask) else 0.0
            X[~mask, col] = median_val
            n_imputed += nan_count

    if n_imputed > 0:
        logger.info(
            "Imputed %d NaN values with column medians (%.2f%% of matrix)",
            n_imputed, 100.0 * n_imputed / X.size,
        )

    if method == "zscore":
        means = X.mean(axis=0)
        stds = X.std(axis=0)
        n_constant = int(np.sum(stds < 1e-12))
        stds[stds < 1e-12] = 1.0  # avoid division by zero for constant features
        X = (X - means) / stds
        stats["means"] = means.tolist()
        stats["stds"] = stds.tolist()
        if n_constant > 0:
            logger.warning(
                "%d feature(s) have near-zero variance and were not scaled.",
                n_constant,
            )
    elif method == "robust":
        medians = np.median(X, axis=0)
        q75 = np.percentile(X, 75, axis=0)
        q25 = np.percentile(X, 25, axis=0)
        iqr = q75 - q25
        iqr[iqr < 1e-12] = 1.0
        X = (X - medians) / iqr
        stats["medians"] = medians.tolist()
        stats["iqr"] = iqr.tolist()
    elif method == "minmax":
        mins = X.min(axis=0)
        maxs = X.max(axis=0)
        ranges = maxs - mins
        ranges[ranges < 1e-12] = 1.0
        X = (X - mins) / ranges
        stats["mins"] = mins.tolist()
        stats["ranges"] = ranges.tolist()

    # Replace any inf with 0
    X[~np.isfinite(X)] = 0.0
    return X, stats


# ---------------------------------------------------------------------------
# k-NN graph + clustering
# ---------------------------------------------------------------------------

def _build_knn_graph(X: np.ndarray, n_neighbors: int) -> Any:
    """Build a k-NN graph and return as igraph Graph or scipy sparse matrix."""
    from sklearn.neighbors import NearestNeighbors

    t0 = time.monotonic()
    logger.info("  Computing %d nearest neighbors for %d points...", n_neighbors, X.shape[0])

    nn = NearestNeighbors(n_neighbors=n_neighbors, metric="euclidean", n_jobs=-1)
    nn.fit(X)
    distances, indices = nn.kneighbors(X)

    knn_elapsed = time.monotonic() - t0
    logger.info("  k-NN search done in %s", _fmt_duration(knn_elapsed))

    n = X.shape[0]
    t0 = time.monotonic()

    if _HAS_LEIDEN:
        logger.info("  Building igraph weighted graph (%d nodes)...", n)
        edges = []
        weights = []
        max_d = distances[:, -1].max() if distances.size > 0 else 1.0
        max_d = max(max_d, 1e-12)
        for i in range(n):
            for j_idx in range(n_neighbors):
                j = indices[i, j_idx]
                if i != j:
                    edges.append((i, j))
                    weights.append(1.0 - distances[i, j_idx] / max_d)
        g = ig.Graph(n=n, edges=edges, directed=True)
        g.es["weight"] = weights
        g = g.as_undirected(mode="collapse", combine_edges={"weight": "max"})
        graph_elapsed = time.monotonic() - t0
        logger.info(
            "  Graph built in %s: %d nodes, %d edges",
            _fmt_duration(graph_elapsed), g.vcount(), g.ecount(),
        )
        return g

    # Fallback: return sparse affinity matrix for sklearn
    logger.info("  Building sparse affinity matrix (%d x %d)...", n, n)
    from scipy.sparse import lil_matrix
    W = lil_matrix((n, n), dtype=np.float64)
    max_d = distances[:, -1].max() if distances.size > 0 else 1.0
    max_d = max(max_d, 1e-12)
    for i in range(n):
        for j_idx in range(n_neighbors):
            j = indices[i, j_idx]
            if i != j:
                w = 1.0 - distances[i, j_idx] / max_d
                W[i, j] = max(W[i, j], w)
                W[j, i] = max(W[j, i], w)
    csr = W.tocsr()
    graph_elapsed = time.monotonic() - t0
    logger.info(
        "  Affinity matrix built in %s: %d nonzero entries",
        _fmt_duration(graph_elapsed), csr.nnz,
    )
    return csr


def _cluster_leiden(graph: Any, resolution: float) -> np.ndarray:
    """Run Leiden community detection."""
    partition = leidenalg.find_partition(
        graph,
        leidenalg.RBConfigurationVertexPartition,
        weights="weight",
        resolution_parameter=resolution,
        seed=42,
    )
    return np.array(partition.membership, dtype=np.int32)


def _cluster_spectral(affinity: Any, n_clusters: int = 20) -> np.ndarray:
    """Fallback: spectral clustering via sklearn."""
    from sklearn.cluster import SpectralClustering

    sc = SpectralClustering(
        n_clusters=min(n_clusters, affinity.shape[0]),
        affinity="precomputed",
        random_state=42,
        n_init=5,
    )
    return sc.fit_predict(affinity).astype(np.int32)


# ---------------------------------------------------------------------------
# Embedding
# ---------------------------------------------------------------------------

def _embed_umap(X: np.ndarray, cfg: ClusteringConfig) -> np.ndarray:
    """UMAP 2D embedding."""
    # Ensure no NaN/inf — older umap-learn triggers sklearn check_array
    # incompatibilities with newer sklearn (force_all_finite renamed)
    X_clean = np.nan_to_num(X, nan=0.0, posinf=0.0, neginf=0.0)
    reducer = umap.UMAP(
        n_components=cfg.umap_n_components,
        n_neighbors=cfg.n_neighbors,
        min_dist=cfg.umap_min_dist,
        metric=cfg.umap_metric,
        random_state=42,
    )
    return reducer.fit_transform(X_clean).astype(np.float64)


def _embed_tsne(X: np.ndarray) -> np.ndarray:
    """Fallback: t-SNE via sklearn."""
    from sklearn.manifold import TSNE

    perp = min(30, max(5, X.shape[0] // 5))
    tsne = TSNE(
        n_components=2,
        perplexity=perp,
        random_state=42,
        init="pca",
        learning_rate="auto",
    )
    return tsne.fit_transform(X).astype(np.float64)


# ---------------------------------------------------------------------------
# Main clustering function
# ---------------------------------------------------------------------------

def run_clustering(
    dataset: BehaviorDataset,
    cfg: ClusteringConfig | None = None,
) -> ClusteringResult:
    """Normalize features, build k-NN graph, cluster, embed.

    Parameters
    ----------
    dataset : BehaviorDataset with (N, 33) features
    cfg : clustering configuration

    Returns
    -------
    ClusteringResult with labels, embedding, normalized features
    """
    if cfg is None:
        cfg = ClusteringConfig()

    n_samples = len(dataset)

    logger.info("=" * 70)
    logger.info("CLUSTERING")
    logger.info("  Samples: %d windows, %d features", n_samples, dataset.n_features)
    logger.info("  Method: %s", "Leiden" if _HAS_LEIDEN else "Spectral (leidenalg not installed)")
    logger.info("  Embedding: %s", "UMAP" if _HAS_UMAP else "t-SNE (umap-learn not installed)")
    logger.info("  k-NN neighbors: %d", cfg.n_neighbors)
    if _HAS_LEIDEN:
        logger.info("  Leiden resolution: %.2f", cfg.leiden_resolution)
    logger.info("  Normalization: %s", cfg.normalize_method)
    logger.info("=" * 70)

    if n_samples == 0:
        logger.warning("No samples to cluster.")
        return ClusteringResult(
            features_normalized=dataset.features,
            cluster_labels=np.array([], dtype=np.int32),
            n_clusters=0,
            embedding_2d=np.empty((0, 2), dtype=np.float64),
            metadata=dataset.metadata,
            feature_names=dataset.feature_names,
            method="none",
            embedding_method="none",
        )

    pipeline_t0 = time.monotonic()

    # 1. Normalize
    logger.info("Step 1/4: Normalizing features (%s)...", cfg.normalize_method)
    t0 = time.monotonic()
    X_norm, norm_stats = _normalize_features(dataset.features, cfg.normalize_method)
    logger.info(
        "  Normalization done in %s. Shape: %s",
        _fmt_duration(time.monotonic() - t0), X_norm.shape,
    )

    # 2. Build k-NN graph
    k = min(cfg.n_neighbors, n_samples - 1)
    logger.info("Step 2/4: Building k-NN graph (k=%d)...", k)
    t0 = time.monotonic()
    graph = _build_knn_graph(X_norm, k)
    logger.info("  k-NN graph done in %s", _fmt_duration(time.monotonic() - t0))

    # 3. Cluster
    logger.info("Step 3/4: Clustering...")
    t0 = time.monotonic()
    if _HAS_LEIDEN:
        logger.info("  Running Leiden (resolution=%.2f)...", cfg.leiden_resolution)
        labels = _cluster_leiden(graph, cfg.leiden_resolution)
        method = "leiden"
    else:
        logger.info("  Running spectral clustering (fallback)...")
        labels = _cluster_spectral(graph)
        method = "spectral"
    cluster_elapsed = time.monotonic() - t0

    n_clusters = int(labels.max() + 1) if len(labels) > 0 else 0
    logger.info(
        "  Found %d clusters via %s in %s",
        n_clusters, method, _fmt_duration(cluster_elapsed),
    )

    # Log cluster sizes
    for c in range(n_clusters):
        size = int(np.sum(labels == c))
        logger.info("    Cluster %d: %d windows (%.1f%%)", c, size, 100.0 * size / n_samples)

    # 4. Embed
    emb_name = "UMAP" if _HAS_UMAP else "t-SNE"
    logger.info("Step 4/4: Computing %s embedding...", emb_name)
    t0 = time.monotonic()
    if _HAS_UMAP:
        try:
            embedding = _embed_umap(X_norm, cfg)
            emb_method = "umap"
        except Exception as e:
            logger.warning("UMAP failed (%s), falling back to t-SNE...", e)
            embedding = _embed_tsne(X_norm)
            emb_method = "tsne"
    else:
        embedding = _embed_tsne(X_norm)
        emb_method = "tsne"
    emb_elapsed = time.monotonic() - t0
    logger.info("  %s done in %s. Shape: %s", emb_name, _fmt_duration(emb_elapsed), embedding.shape)

    total_elapsed = time.monotonic() - pipeline_t0
    logger.info("=" * 70)
    logger.info("CLUSTERING COMPLETE")
    logger.info("  %d clusters from %d windows", n_clusters, n_samples)
    logger.info("  Method: %s | Embedding: %s", method, emb_method)
    logger.info("  Total clustering time: %s", _fmt_duration(total_elapsed))
    logger.info("=" * 70)

    return ClusteringResult(
        features_normalized=X_norm,
        cluster_labels=labels,
        n_clusters=n_clusters,
        embedding_2d=embedding,
        metadata=dataset.metadata,
        feature_names=dataset.feature_names,
        method=method,
        embedding_method=emb_method,
    )


# ---------------------------------------------------------------------------
# Pipeline class
# ---------------------------------------------------------------------------

class ClusteringPipeline:
    """End-to-end clustering pipeline for a project folder.

    Usage::

        pipe = ClusteringPipeline(project_dir)
        result = pipe.run()
    """

    def __init__(
        self,
        project_dir: Path,
        *,
        cfg: ClusteringConfig | None = None,
        show_progress: bool = True,
        skip_mask_verification: bool = False,
    ):
        self.project_dir = Path(project_dir).resolve()
        self.cfg = cfg or ClusteringConfig()
        self.show_progress = show_progress
        self.skip_mask_verification = skip_mask_verification

    @property
    def output_dir(self) -> Path:
        return self.project_dir / "analysis_of_tracking_data" / "clustering"

    def run(self) -> ClusteringResult:
        out = self.output_dir
        out.mkdir(parents=True, exist_ok=True)
        overall_t0 = time.monotonic()

        logger.info("#" * 70)
        logger.info("# BEHAVIORAL SEQUENCE CLUSTERING PIPELINE")
        logger.info("#" * 70)
        logger.info("Project dir: %s", self.project_dir)
        logger.info("Output dir:  %s", out)
        logger.info("Config: window=%d stride=%d k=%d resolution=%.2f normalize=%s",
                     self.cfg.window_size, self.cfg.stride,
                     self.cfg.n_neighbors, self.cfg.leiden_resolution,
                     self.cfg.normalize_method)
        logger.info("")

        # 1. Feature extraction
        logger.info("PHASE 1: Feature extraction")
        logger.info("-" * 40)
        t0 = time.monotonic()
        dataset = extract_all_features(
            self.project_dir, self.cfg,
            progress=self.show_progress,
            skip_mask_verification=self.skip_mask_verification,
        )
        phase1_time = time.monotonic() - t0
        logger.info("Phase 1 done in %s: %d windows extracted.", _fmt_duration(phase1_time), len(dataset))
        logger.info("")

        if len(dataset) == 0:
            logger.warning("No features extracted. Writing empty summary.")
            self._write_summary(out, dataset, None)
            return ClusteringResult(
                features_normalized=dataset.features,
                cluster_labels=np.array([], dtype=np.int32),
                n_clusters=0,
                embedding_2d=np.empty((0, 2), dtype=np.float64),
                metadata=dataset.metadata,
                feature_names=dataset.feature_names,
                method="none",
                embedding_method="none",
            )

        # 2. Clustering
        logger.info("PHASE 2: Clustering + embedding")
        logger.info("-" * 40)
        t0 = time.monotonic()
        result = run_clustering(dataset, self.cfg)
        phase2_time = time.monotonic() - t0
        logger.info("Phase 2 done in %s.", _fmt_duration(phase2_time))
        logger.info("")

        # 3. Downstream comparison
        logger.info("PHASE 3: Housing condition comparison")
        logger.info("-" * 40)
        t0 = time.monotonic()
        from .comparison import compare_housing_conditions
        comparison = compare_housing_conditions(result)
        phase3_time = time.monotonic() - t0
        logger.info("Phase 3 done in %s.", _fmt_duration(phase3_time))
        logger.info("")

        # 4. Plots
        logger.info("PHASE 4: Generating plots")
        logger.info("-" * 40)
        t0 = time.monotonic()
        from .plots import generate_all_plots
        plots_dir = out / "plots"
        plots_dir.mkdir(parents=True, exist_ok=True)
        generate_all_plots(result, comparison, plots_dir)
        phase4_time = time.monotonic() - t0
        logger.info("Phase 4 done in %s.", _fmt_duration(phase4_time))
        logger.info("")

        # 5. Save results
        logger.info("PHASE 5: Saving results")
        logger.info("-" * 40)
        results_dir = out / "results"
        results_dir.mkdir(parents=True, exist_ok=True)
        self._save_results(results_dir, result, comparison)
        self._write_summary(out, dataset, result, comparison)
        self._write_excel_report(out, result, comparison)

        overall_elapsed = time.monotonic() - overall_t0
        logger.info("")
        logger.info("#" * 70)
        logger.info("# PIPELINE COMPLETE")
        logger.info("#   Clusters: %d", result.n_clusters)
        logger.info("#   Windows:  %d", len(result.metadata))
        logger.info("#   Time breakdown:")
        logger.info("#     Feature extraction: %s", _fmt_duration(phase1_time))
        logger.info("#     Clustering:         %s", _fmt_duration(phase2_time))
        logger.info("#     Comparison:         %s", _fmt_duration(phase3_time))
        logger.info("#     Plots:              %s", _fmt_duration(phase4_time))
        logger.info("#     Total:              %s", _fmt_duration(overall_elapsed))
        logger.info("#   Output: %s", out)
        logger.info("#" * 70)

        return result

    def _save_results(
        self,
        results_dir: Path,
        result: ClusteringResult,
        comparison: dict[str, Any],
    ) -> None:
        # Cluster assignments
        assignments = []
        for i, m in enumerate(result.metadata):
            assignments.append({
                "window_index": i,
                "cluster": int(result.cluster_labels[i]),
                "video_id": m.video_id,
                "video_name": m.video_name,
                "start_frame": m.start_frame,
                "object_a_name": m.object_a_name,
                "object_b_name": m.object_b_name,
                "interaction_type": m.interaction_type,
                "umap_x": float(result.embedding_2d[i, 0]),
                "umap_y": float(result.embedding_2d[i, 1]),
            })
        p = results_dir / "cluster_assignments.json"
        p.write_text(json.dumps(assignments, indent=2))
        logger.info("  Wrote %s (%s)", p.name, _fmt_size(p.stat().st_size))

        # CSV: cluster assignments + features (one row per window)
        csv_path = results_dir / "cluster_assignments.csv"
        with open(csv_path, "w", newline="") as f:
            meta_cols = [
                "window_index", "cluster", "video_id", "video_name",
                "start_frame", "object_a_name", "object_b_name",
                "interaction_type", "umap_x", "umap_y",
            ]
            writer = csv.DictWriter(f, fieldnames=meta_cols + list(result.feature_names))
            writer.writeheader()
            for i, row in enumerate(assignments):
                for fi, fname in enumerate(result.feature_names):
                    row[fname] = float(result.features_normalized[i, fi])
                writer.writerow(row)
        logger.info("  Wrote %s (%s)", csv_path.name, _fmt_size(csv_path.stat().st_size))

        # Cluster summary (centroids + sizes)
        cluster_summary = {}
        for c in range(result.n_clusters):
            mask = result.cluster_labels == c
            centroid = result.features_normalized[mask].mean(axis=0)
            cluster_summary[str(c)] = {
                "size": int(np.sum(mask)),
                "fraction": float(np.sum(mask)) / len(result.cluster_labels),
                "feature_centroid": {
                    name: float(centroid[i])
                    for i, name in enumerate(result.feature_names)
                },
            }
        p = results_dir / "cluster_summary.json"
        p.write_text(json.dumps(cluster_summary, indent=2))
        logger.info("  Wrote %s (%s)", p.name, _fmt_size(p.stat().st_size))

        # CSV: cluster summary
        csv_path = results_dir / "cluster_summary.csv"
        with open(csv_path, "w", newline="") as f:
            cols = ["cluster", "size", "fraction"] + list(result.feature_names)
            writer = csv.writer(f)
            writer.writerow(cols)
            for c in range(result.n_clusters):
                cs = cluster_summary[str(c)]
                row = [c, cs["size"], cs["fraction"]]
                row += [cs["feature_centroid"][name] for name in result.feature_names]
                writer.writerow(row)
        logger.info("  Wrote %s (%s)", csv_path.name, _fmt_size(csv_path.stat().st_size))

        # Housing comparison
        p = results_dir / "housing_comparison.json"
        p.write_text(json.dumps(comparison, indent=2, default=str))
        logger.info("  Wrote %s (%s)", p.name, _fmt_size(p.stat().st_size))

        # CSV: per-cluster enrichment
        csv_path = results_dir / "housing_enrichment_per_cluster.csv"
        per_cluster = comparison.get("per_cluster", {})
        with open(csv_path, "w", newline="") as f:
            writer = csv.writer(f)
            writer.writerow([
                "cluster", "size", "fraction",
                "isolated_in_cluster", "group_in_cluster",
                "isolated_outside", "group_outside",
                "odds_ratio", "p_value",
            ])
            for c_str in sorted(per_cluster, key=int):
                pc = per_cluster[c_str]
                enr = pc["isolated_enrichment"]
                writer.writerow([
                    c_str, pc["size"], pc["fraction"],
                    enr["isolated_in_cluster"], enr["group_in_cluster"],
                    enr["isolated_outside"], enr["group_outside"],
                    enr["odds_ratio"], enr["p_value"],
                ])
        logger.info("  Wrote %s (%s)", csv_path.name, _fmt_size(csv_path.stat().st_size))

        # CSV: per-feature Mann-Whitney U tests
        per_feat = comparison.get("per_feature_tests", {})
        if per_feat:
            csv_path = results_dir / "feature_tests.csv"
            with open(csv_path, "w", newline="") as f:
                writer = csv.writer(f)
                writer.writerow([
                    "feature", "comparison", "U_statistic", "p_value",
                    "mean_a", "mean_b", "effect_direction",
                ])
                for fname in result.feature_names:
                    if fname in per_feat:
                        ft = per_feat[fname]
                        writer.writerow([
                            fname, ft["comparison"], ft["U_statistic"],
                            ft["p_value"], ft["mean_a"], ft["mean_b"],
                            ft["effect_direction"],
                        ])
            logger.info("  Wrote %s (%s)", csv_path.name, _fmt_size(csv_path.stat().st_size))

        # CSV: per-feature enrichment (isolated vs group-only)
        feat_enr = comparison.get("feature_enrichment", [])
        if feat_enr:
            csv_path = results_dir / "feature_enrichment.csv"
            with open(csv_path, "w", newline="") as f:
                cols = [
                    "feature", "mean_isolated", "mean_group",
                    "std_isolated", "std_group", "cohens_d",
                    "log2_fold_change", "odds_ratio", "fisher_p_value",
                    "mannwhitney_U", "mannwhitney_p_value",
                    "n_isolated", "n_group", "direction",
                ]
                writer = csv.DictWriter(f, fieldnames=cols)
                writer.writeheader()
                for row in feat_enr:
                    writer.writerow({k: row[k] for k in cols})
            logger.info("  Wrote %s (%s)", csv_path.name, _fmt_size(csv_path.stat().st_size))

        # Embedding coordinates
        p = results_dir / "embedding.npz"
        np.savez_compressed(
            p,
            embedding=result.embedding_2d,
            labels=result.cluster_labels,
        )
        logger.info("  Wrote %s (%s)", p.name, _fmt_size(p.stat().st_size))

    def _write_summary(
        self,
        out_dir: Path,
        dataset: BehaviorDataset,
        result: ClusteringResult | None,
        comparison: dict[str, Any] | None = None,
    ) -> None:
        summary: dict[str, Any] = {
            "project_dir": str(self.project_dir),
            "config": {
                "window_size": self.cfg.window_size,
                "stride": self.cfg.stride,
                "n_neighbors": self.cfg.n_neighbors,
                "leiden_resolution": self.cfg.leiden_resolution,
                "normalize_method": self.cfg.normalize_method,
            },
            "n_windows": len(dataset),
            "n_features": dataset.n_features,
        }
        if result is not None:
            summary["n_clusters"] = result.n_clusters
            summary["clustering_method"] = result.method
            summary["embedding_method"] = result.embedding_method

            # Cluster size distribution
            sizes = []
            for c in range(result.n_clusters):
                sizes.append(int(np.sum(result.cluster_labels == c)))
            summary["cluster_sizes"] = sizes

        if comparison is not None:
            summary["housing_comparison_summary"] = {
                k: v for k, v in comparison.items()
                if k in ("overall_enrichment", "interaction_type_distribution")
            }

        p = out_dir / "summary.json"
        p.write_text(json.dumps(summary, indent=2, default=str))
        logger.info("  Wrote %s", p.name)

    def _write_excel_report(
        self,
        out_dir: Path,
        result: ClusteringResult,
        comparison: dict[str, Any],
    ) -> None:
        """Write a single Excel workbook with all results and embedded figures."""
        try:
            from openpyxl import Workbook
            from openpyxl.drawing.image import Image as XlImage
            from openpyxl.utils import get_column_letter
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            logger.warning("openpyxl not installed — skipping Excel report. "
                           "Install with: pip install openpyxl")
            return

        logger.info("  Writing Excel report...")
        t0 = time.monotonic()
        wb = Workbook()

        header_font = Font(bold=True, size=11)
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font_white = Font(bold=True, size=11, color="FFFFFF")
        thin_border = Border(
            bottom=Side(style="thin", color="CCCCCC"),
        )
        pct_fmt = "0.0%"
        sci_fmt = "0.00E+00"
        float_fmt = "0.000"

        def _style_header(ws, n_cols: int) -> None:
            for col in range(1, n_cols + 1):
                cell = ws.cell(row=1, column=col)
                cell.font = header_font_white
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", wrap_text=True)

        def _auto_width(ws, max_width: int = 30) -> None:
            for col_cells in ws.columns:
                length = max(
                    len(str(cell.value or "")) for cell in col_cells
                )
                col_letter = get_column_letter(col_cells[0].column)
                ws.column_dimensions[col_letter].width = min(length + 3, max_width)

        # ---- Sheet 1: Summary ----
        ws = wb.active
        ws.title = "Summary"
        summary_rows = [
            ("Parameter", "Value"),
            ("Project", self.project_dir.name),
            ("Total windows", len(result.metadata)),
            ("Features per window", len(result.feature_names)),
            ("Clusters found", result.n_clusters),
            ("Clustering method", result.method),
            ("Embedding method", result.embedding_method),
            ("Window size (frames)", self.cfg.window_size),
            ("Stride (frames)", self.cfg.stride),
            ("k-NN neighbors", self.cfg.n_neighbors),
            ("Leiden resolution", self.cfg.leiden_resolution),
            ("Normalization", self.cfg.normalize_method),
        ]
        overall = comparison.get("overall_enrichment", {})
        if overall:
            summary_rows.append(("Windows with isolated mouse", overall.get("total_windows_with_isolated", "")))
            summary_rows.append(("Windows group-only", overall.get("total_windows_group_only", "")))
        for r, (k, v) in enumerate(summary_rows, 1):
            ws.cell(row=r, column=1, value=k)
            ws.cell(row=r, column=2, value=v)
        ws.cell(row=1, column=1).font = header_font
        ws.cell(row=1, column=2).font = header_font
        _auto_width(ws)

        # ---- Sheet 2: Cluster Summary ----
        ws = wb.create_sheet("Cluster Summary")
        cols = ["Cluster", "Size", "Fraction"] + list(result.feature_names)
        for c_idx, col_name in enumerate(cols, 1):
            ws.cell(row=1, column=c_idx, value=col_name)
        _style_header(ws, len(cols))
        for c in range(result.n_clusters):
            mask = result.cluster_labels == c
            centroid = result.features_normalized[mask].mean(axis=0)
            row = c + 2
            ws.cell(row=row, column=1, value=c)
            ws.cell(row=row, column=2, value=int(np.sum(mask)))
            ws.cell(row=row, column=3, value=float(np.sum(mask)) / len(result.cluster_labels))
            ws[f"C{row}"].number_format = pct_fmt
            for fi, fname in enumerate(result.feature_names):
                ws.cell(row=row, column=4 + fi, value=float(centroid[fi]))
                ws.cell(row=row, column=4 + fi).number_format = float_fmt
        _auto_width(ws)

        # ---- Sheet 3: Cluster Assignments (capped at 100k rows for Excel) ----
        ws = wb.create_sheet("Cluster Assignments")
        meta_cols = [
            "Window", "Cluster", "Video ID", "Video Name",
            "Start Frame", "Object A", "Object B",
            "Interaction Type", "UMAP X", "UMAP Y",
        ]
        all_cols = meta_cols + list(result.feature_names)
        for c_idx, col_name in enumerate(all_cols, 1):
            ws.cell(row=1, column=c_idx, value=col_name)
        _style_header(ws, len(all_cols))
        n_rows = min(len(result.metadata), 100000)
        if n_rows < len(result.metadata):
            logger.info("    Cluster Assignments sheet capped at %d rows (of %d)", n_rows, len(result.metadata))
        for i in range(n_rows):
            m = result.metadata[i]
            row = i + 2
            ws.cell(row=row, column=1, value=i)
            ws.cell(row=row, column=2, value=int(result.cluster_labels[i]))
            ws.cell(row=row, column=3, value=m.video_id)
            ws.cell(row=row, column=4, value=m.video_name)
            ws.cell(row=row, column=5, value=m.start_frame)
            ws.cell(row=row, column=6, value=m.object_a_name)
            ws.cell(row=row, column=7, value=m.object_b_name)
            ws.cell(row=row, column=8, value=m.interaction_type)
            ws.cell(row=row, column=9, value=float(result.embedding_2d[i, 0]))
            ws.cell(row=row, column=10, value=float(result.embedding_2d[i, 1]))
            for fi in range(len(result.feature_names)):
                ws.cell(row=row, column=11 + fi, value=float(result.features_normalized[i, fi]))
        _auto_width(ws, max_width=18)

        # ---- Sheet 4: Housing Enrichment per Cluster ----
        ws = wb.create_sheet("Cluster Enrichment")
        per_cluster = comparison.get("per_cluster", {})
        enr_cols = [
            "Cluster", "Size", "Fraction",
            "Isolated in Cluster", "Group in Cluster",
            "Isolated Outside", "Group Outside",
            "Odds Ratio", "p-value", "Significant",
        ]
        for c_idx, col_name in enumerate(enr_cols, 1):
            ws.cell(row=1, column=c_idx, value=col_name)
        _style_header(ws, len(enr_cols))
        for ri, c_str in enumerate(sorted(per_cluster, key=int), 2):
            pc = per_cluster[c_str]
            enr = pc["isolated_enrichment"]
            p_val = enr["p_value"]
            sig = "***" if (np.isfinite(p_val) and p_val < 0.001) else \
                  "**" if (np.isfinite(p_val) and p_val < 0.01) else \
                  "*" if (np.isfinite(p_val) and p_val < 0.05) else ""
            ws.cell(row=ri, column=1, value=int(c_str))
            ws.cell(row=ri, column=2, value=pc["size"])
            ws.cell(row=ri, column=3, value=pc["fraction"])
            ws[f"C{ri}"].number_format = pct_fmt
            ws.cell(row=ri, column=4, value=enr["isolated_in_cluster"])
            ws.cell(row=ri, column=5, value=enr["group_in_cluster"])
            ws.cell(row=ri, column=6, value=enr["isolated_outside"])
            ws.cell(row=ri, column=7, value=enr["group_outside"])
            ws.cell(row=ri, column=8, value=enr["odds_ratio"])
            ws[f"H{ri}"].number_format = float_fmt
            ws.cell(row=ri, column=9, value=p_val)
            ws[f"I{ri}"].number_format = sci_fmt
            ws.cell(row=ri, column=10, value=sig)
        _auto_width(ws)

        # ---- Sheet 5: Feature Enrichment (isolated vs group-only) ----
        feat_enr = comparison.get("feature_enrichment", [])
        if feat_enr:
            ws = wb.create_sheet("Feature Enrichment")
            fe_cols = [
                "Feature", "Mean (Isolated)", "Mean (Group)",
                "Std (Isolated)", "Std (Group)",
                "Cohen's d", "Log2 FC", "Odds Ratio",
                "Fisher p-value", "Mann-Whitney U", "MW p-value",
                "n Isolated", "n Group", "Direction", "Significant",
            ]
            for c_idx, col_name in enumerate(fe_cols, 1):
                ws.cell(row=1, column=c_idx, value=col_name)
            _style_header(ws, len(fe_cols))
            for ri, fe in enumerate(feat_enr, 2):
                p_val = fe["mannwhitney_p_value"]
                sig = "***" if (np.isfinite(p_val) and p_val < 0.001) else \
                      "**" if (np.isfinite(p_val) and p_val < 0.01) else \
                      "*" if (np.isfinite(p_val) and p_val < 0.05) else ""
                ws.cell(row=ri, column=1, value=fe["feature"])
                ws.cell(row=ri, column=2, value=fe["mean_isolated"])
                ws.cell(row=ri, column=3, value=fe["mean_group"])
                ws.cell(row=ri, column=4, value=fe["std_isolated"])
                ws.cell(row=ri, column=5, value=fe["std_group"])
                ws.cell(row=ri, column=6, value=fe["cohens_d"])
                ws.cell(row=ri, column=7, value=fe["log2_fold_change"])
                ws.cell(row=ri, column=8, value=fe["odds_ratio"])
                ws.cell(row=ri, column=9, value=fe["fisher_p_value"])
                ws.cell(row=ri, column=10, value=fe["mannwhitney_U"])
                ws.cell(row=ri, column=11, value=fe["mannwhitney_p_value"])
                ws.cell(row=ri, column=12, value=fe["n_isolated"])
                ws.cell(row=ri, column=13, value=fe["n_group"])
                ws.cell(row=ri, column=14, value=fe["direction"])
                ws.cell(row=ri, column=15, value=sig)
                for col in (2, 3, 4, 5, 6, 7, 8):
                    ws.cell(row=ri, column=col).number_format = float_fmt
                for col in (9, 11):
                    ws.cell(row=ri, column=col).number_format = sci_fmt
            _auto_width(ws)

        # ---- Sheet 6: Feature Tests (interaction type comparison) ----
        per_feat = comparison.get("per_feature_tests", {})
        if per_feat:
            ws = wb.create_sheet("Feature Tests")
            ft_cols = [
                "Feature", "Comparison", "U Statistic", "p-value",
                "Mean A", "Mean B", "Direction", "Significant",
            ]
            for c_idx, col_name in enumerate(ft_cols, 1):
                ws.cell(row=1, column=c_idx, value=col_name)
            _style_header(ws, len(ft_cols))
            ri = 2
            for fname in result.feature_names:
                if fname not in per_feat:
                    continue
                ft = per_feat[fname]
                p_val = ft["p_value"]
                sig = "***" if (np.isfinite(p_val) and p_val < 0.001) else \
                      "**" if (np.isfinite(p_val) and p_val < 0.01) else \
                      "*" if (np.isfinite(p_val) and p_val < 0.05) else ""
                ws.cell(row=ri, column=1, value=fname)
                ws.cell(row=ri, column=2, value=ft["comparison"])
                ws.cell(row=ri, column=3, value=ft["U_statistic"])
                ws.cell(row=ri, column=4, value=ft["p_value"])
                ws[f"D{ri}"].number_format = sci_fmt
                ws.cell(row=ri, column=5, value=ft["mean_a"])
                ws[f"E{ri}"].number_format = float_fmt
                ws.cell(row=ri, column=6, value=ft["mean_b"])
                ws[f"F{ri}"].number_format = float_fmt
                ws.cell(row=ri, column=7, value=ft["effect_direction"])
                ws.cell(row=ri, column=8, value=sig)
                ri += 1
            _auto_width(ws)

        # ---- Sheet 7+: Figures ----
        plots_dir = out_dir / "plots"
        plot_files = [
            ("UMAP Clusters", "umap_by_cluster.png"),
            ("UMAP Housing", "umap_by_housing.png"),
            ("Cluster Composition", "cluster_composition.png"),
            ("Feature Heatmap", "feature_heatmap.png"),
            ("Cluster Enrichment Plot", "enrichment_bars.png"),
            ("Feature Enrichment Plot", "feature_enrichment.png"),
            ("Feature Violins", "feature_violins.png"),
        ]
        for sheet_name, filename in plot_files:
            img_path = plots_dir / filename
            if not img_path.is_file():
                continue
            ws = wb.create_sheet(sheet_name)
            try:
                img = XlImage(str(img_path))
                # Scale to reasonable size in Excel
                img.width = min(img.width, 900)
                img.height = min(img.height, 700)
                ws.add_image(img, "A1")
            except Exception as e:
                ws.cell(row=1, column=1, value=f"Could not embed {filename}: {e}")

        # Add ethogram images (first 10 to keep file size reasonable)
        ethogram_dir = plots_dir / "ethograms"
        if ethogram_dir.is_dir():
            ethogram_files = sorted(ethogram_dir.glob("ethogram_*.png"))[:10]
            if ethogram_files:
                ws = wb.create_sheet("Ethograms")
                current_row = 1
                for img_path in ethogram_files:
                    try:
                        img = XlImage(str(img_path))
                        img.width = min(img.width, 1100)
                        img.height = min(img.height, 200)
                        ws.add_image(img, f"A{current_row}")
                        current_row += 12  # ~12 rows per ethogram strip
                    except Exception:
                        pass

        # Save
        xlsx_path = out_dir / "clustering_report.xlsx"
        wb.save(str(xlsx_path))
        logger.info("  Wrote %s (%s)", xlsx_path.name, _fmt_size(xlsx_path.stat().st_size))
        elapsed = time.monotonic() - t0
        logger.info("  Excel report done in %s", _fmt_duration(elapsed))


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main(argv: list[str] | None = None) -> int:
    # Configure logging IMMEDIATELY so all messages are visible from the start
    # (parse args first just to get verbosity level)
    p = argparse.ArgumentParser(
        description="Behavioral sequence clustering for SAM3 Web Tracker projects.",
    )
    p.add_argument("--project-dir", type=Path, help="Folder containing config.json")
    p.add_argument("--project-id", type=str, help="Short project uuid")
    p.add_argument("--resolution", type=float, default=1.0, help="Leiden resolution (default: 1.0)")
    p.add_argument("--n-neighbors", type=int, default=30, help="k-NN neighbors (default: 30)")
    p.add_argument("--window-size", type=int, default=90, help="Window size in frames (default: 90)")
    p.add_argument("--stride", type=int, default=30, help="Stride in frames (default: 30)")
    p.add_argument("-v", "--verbose", action="store_true", help="DEBUG logging")
    p.add_argument("-q", "--quiet", action="store_true", help="WARNING only")
    p.add_argument("--no-progress", action="store_true", help="Disable tqdm bars")
    p.add_argument(
        "--skip-mask-verification",
        action="store_true",
        help="Skip per-frame mask existence check (faster startup, relies on propagation_complete flag)",
    )
    args = p.parse_args(argv)

    if args.verbose and args.quiet:
        print("Cannot use both --verbose and --quiet", file=sys.stderr)
        return 2

    # Set up logging before ANYTHING else.
    # Use stdout with line-buffered flush so output appears immediately
    # even under conda run (which buffers stderr).
    if args.quiet:
        level = logging.WARNING
    elif args.verbose:
        level = logging.DEBUG
    else:
        level = logging.INFO

    # Remove any existing handlers, install a flush-on-every-write handler on stdout
    root = logging.getLogger()
    root.handlers.clear()
    handler = logging.StreamHandler(sys.stdout)
    handler.setFormatter(logging.Formatter(
        "%(asctime)s | %(levelname)-8s | %(message)s",
        datefmt="%H:%M:%S",
    ))
    handler.flush = lambda: sys.stdout.flush()  # force flush after every record
    root.addHandler(handler)
    root.setLevel(level)
    logging.getLogger("matplotlib").setLevel(logging.WARNING)
    logging.getLogger("PIL").setLevel(logging.WARNING)

    # Also make sure stdout is unbuffered for print() calls
    sys.stdout.reconfigure(line_buffering=True)

    logger.info("Logging configured. Parsing arguments...")

    show_progress = not args.no_progress and not args.quiet

    if args.project_dir:
        proj = Path(args.project_dir).resolve()
        if not (proj / "config.json").is_file():
            print("No config.json in", proj, file=sys.stderr)
            return 1
    elif args.project_id:
        logger.info("Looking up project id: %s", args.project_id)
        found = pathutil.find_project_dir(args.project_id.strip())
        if found is None:
            print("Project id not found:", args.project_id, file=sys.stderr)
            return 1
        proj = found
    else:
        p.print_help()
        return 2

    logger.info("Project dir: %s", proj)
    logger.info("Args: window=%d stride=%d k=%d resolution=%.2f",
                args.window_size, args.stride, args.n_neighbors, args.resolution)

    cfg = ClusteringConfig(
        window_size=args.window_size,
        stride=args.stride,
        n_neighbors=args.n_neighbors,
        leiden_resolution=args.resolution,
    )

    pipe = ClusteringPipeline(
        proj, cfg=cfg, show_progress=show_progress,
        skip_mask_verification=args.skip_mask_verification,
    )
    result = pipe.run()

    if not args.quiet:
        print(f"Done: {result.n_clusters} clusters from {len(result.metadata)} windows")
        print(f"Output: {pipe.output_dir}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
