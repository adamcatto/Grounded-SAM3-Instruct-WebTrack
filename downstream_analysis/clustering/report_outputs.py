"""Summary JSON and Excel report writers for clustering pipelines."""

from __future__ import annotations

import json
import logging
import time
from pathlib import Path
from typing import Any

import numpy as np

from .clustering_pipeline import ClusteringResult, _fmt_duration, _fmt_size
from .config import ClusteringConfig
logger = logging.getLogger(__name__)


def write_summary_json(out_dir: Path, summary: dict[str, Any]) -> None:
    """Write ``summary.json`` at the clustering output root."""
    out_dir.mkdir(parents=True, exist_ok=True)
    p = out_dir / "summary.json"
    p.write_text(json.dumps(summary, indent=2, default=str))
    logger.info("  Wrote %s", p.name)


def build_single_project_summary(
    project_dir: Path,
    cfg: ClusteringConfig,
    dataset: Any,
    result: ClusteringResult | None,
    comparison: dict[str, Any] | None,
) -> dict[str, Any]:
    summary: dict[str, Any] = {
        "pipeline": "single_project",
        "project_dir": str(project_dir),
        "config": {
            "window_size": cfg.window_size,
            "stride": cfg.stride,
            "n_neighbors": cfg.n_neighbors,
            "leiden_resolution": cfg.leiden_resolution,
            "normalize_method": cfg.normalize_method,
        },
        "n_windows": len(dataset),
        "n_features": dataset.n_features,
    }
    if result is not None:
        summary["n_clusters"] = result.n_clusters
        summary["clustering_method"] = result.method
        summary["embedding_method"] = result.embedding_method
        summary["cluster_sizes"] = [
            int(np.sum(result.cluster_labels == c)) for c in range(result.n_clusters)
        ]
    if comparison is not None:
        summary["housing_comparison_summary"] = {
            k: v for k, v in comparison.items()
            if k in ("overall_enrichment", "interaction_type_distribution")
        }
    return summary


def build_multi_project_summary(
    output_dir: Path,
    project_dirs: list[Path],
    cfg: ClusteringConfig,
    result: ClusteringResult,
    comparison: dict[str, Any],
    *,
    batch_correction_method: str,
    locomotion_dir: Path | None = None,
) -> dict[str, Any]:
    condition_groups = comparison.get("condition_groups", {})
    return {
        "pipeline": "multi_project",
        "output_dir": str(output_dir),
        "project_dirs": [str(d) for d in project_dirs],
        "locomotion_dir": str(locomotion_dir) if locomotion_dir else None,
        "batch_correction_method": batch_correction_method,
        "config": {
            "window_size": cfg.window_size,
            "stride": cfg.stride,
            "n_neighbors": cfg.n_neighbors,
            "leiden_resolution": cfg.leiden_resolution,
            "normalize_method": cfg.normalize_method,
        },
        "n_windows": len(result.metadata),
        "n_features": len(result.feature_names),
        "n_clusters": result.n_clusters,
        "clustering_method": result.method,
        "embedding_method": result.embedding_method,
        "cluster_sizes": [
            int(np.sum(result.cluster_labels == c)) for c in range(result.n_clusters)
        ],
        "condition_group_counts": {
            g: info.get("n_windows", 0) for g, info in condition_groups.items()
        },
        "n_pairwise_condition_comparisons": len(
            comparison.get("pairwise_condition_groups", {})
        ),
        "n_pairwise_role_comparisons": len(
            comparison.get("pairwise_role_conditions", {})
        ),
        "n_transition_groups": len(comparison.get("transition_matrices", {})),
        "n_differential_transition_pairs": len(
            comparison.get("differential_transition_matrices", {})
        ),
    }


def _excel_styles():
    try:
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    except ImportError:
        return None
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font_white = Font(bold=True, size=11, color="FFFFFF")
    thin_border = Border(bottom=Side(style="thin", color="CCCCCC"))
    return {
        "header_font": header_font,
        "header_fill": header_fill,
        "header_font_white": header_font_white,
        "thin_border": thin_border,
        "pct_fmt": "0.0%",
        "sci_fmt": "0.00E+00",
        "float_fmt": "0.000",
    }


def _style_header(ws, n_cols: int, styles: dict) -> None:
    from openpyxl.styles import Alignment
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = styles["header_font_white"]
        cell.fill = styles["header_fill"]
        cell.alignment = Alignment(horizontal="center", wrap_text=True)


def _auto_width(ws, max_width: int = 30) -> None:
    from openpyxl.utils import get_column_letter
    for col_cells in ws.columns:
        length = max(len(str(cell.value or "")) for cell in col_cells)
        col_letter = get_column_letter(col_cells[0].column)
        ws.column_dimensions[col_letter].width = min(length + 3, max_width)


def _sig_marker(p_val: float) -> str:
    if not np.isfinite(p_val):
        return ""
    if p_val < 0.001:
        return "***"
    if p_val < 0.01:
        return "**"
    if p_val < 0.05:
        return "*"
    return ""


def _add_cluster_sheets(wb, result: ClusteringResult, cfg: ClusteringConfig, styles: dict) -> None:
    pct_fmt = styles["pct_fmt"]
    float_fmt = styles["float_fmt"]

    ws = wb.create_sheet("Cluster Summary")
    cols = ["Cluster", "Size", "Fraction"] + list(result.feature_names)
    for c_idx, col_name in enumerate(cols, 1):
        ws.cell(row=1, column=c_idx, value=col_name)
    _style_header(ws, len(cols), styles)
    for c in range(result.n_clusters):
        mask = result.cluster_labels == c
        centroid = result.features_normalized[mask].mean(axis=0)
        row = c + 2
        ws.cell(row=row, column=1, value=c)
        ws.cell(row=row, column=2, value=int(np.sum(mask)))
        ws.cell(row=row, column=3, value=float(np.sum(mask)) / len(result.cluster_labels))
        ws[f"C{row}"].number_format = pct_fmt
        for fi, fname in enumerate(result.feature_names):
            ws.cell(row=row, column=4 + fi, value=float(centroid[fi]))
            ws.cell(row=row, column=4 + fi).number_format = float_fmt
    _auto_width(ws)

    ws = wb.create_sheet("Cluster Assignments")
    meta_cols = [
        "Window", "Cluster", "Video ID", "Video Name", "Start Frame",
        "Experiment", "Session", "Batch ID",
        "Mouse A ID", "Mouse B ID", "Mouse A Role", "Mouse B Role",
        "Mouse A Housing", "Mouse B Housing",
        "Object A", "Object B", "Interaction Type", "Camera View",
        "UMAP X", "UMAP Y",
    ]
    all_cols = meta_cols + list(result.feature_names)
    for c_idx, col_name in enumerate(all_cols, 1):
        ws.cell(row=1, column=c_idx, value=col_name)
    _style_header(ws, len(all_cols), styles)
    n_rows = min(len(result.metadata), 100000)
    for i in range(n_rows):
        m = result.metadata[i]
        row = i + 2
        ws.cell(row=row, column=1, value=i)
        ws.cell(row=row, column=2, value=int(result.cluster_labels[i]))
        ws.cell(row=row, column=3, value=m.video_id)
        ws.cell(row=row, column=4, value=m.video_name)
        ws.cell(row=row, column=5, value=m.start_frame)
        ws.cell(row=row, column=6, value=getattr(m, "experiment_name", "") or "")
        ws.cell(row=row, column=7, value=getattr(m, "session", "") or "")
        ws.cell(row=row, column=8, value=getattr(m, "batch_id", "") or "")
        ws.cell(row=row, column=9, value=getattr(m, "mouse_a_id", "") or "")
        ws.cell(row=row, column=10, value=getattr(m, "mouse_b_id", "") or "")
        ws.cell(row=row, column=11, value=getattr(m, "mouse_a_role", "") or "")
        ws.cell(row=row, column=12, value=getattr(m, "mouse_b_role", "") or "")
        ws.cell(row=row, column=13, value=getattr(m, "mouse_a_housing", "") or "")
        ws.cell(row=row, column=14, value=getattr(m, "mouse_b_housing", "") or "")
        ws.cell(row=row, column=15, value=m.object_a_name)
        ws.cell(row=row, column=16, value=m.object_b_name)
        ws.cell(row=row, column=17, value=m.interaction_type)
        ws.cell(row=row, column=18, value=getattr(m, "camera_view", "") or "")
        ws.cell(row=row, column=19, value=float(result.embedding_2d[i, 0]))
        ws.cell(row=row, column=20, value=float(result.embedding_2d[i, 1]))
        for fi in range(len(result.feature_names)):
            ws.cell(row=row, column=21 + fi, value=float(result.features_normalized[i, fi]))
    _auto_width(ws, max_width=18)


def _embed_plot_sheets(wb, plots_dir: Path, plot_files: list[tuple[str, str]]) -> None:
    try:
        from openpyxl.drawing.image import Image as XlImage
    except ImportError:
        return
    for sheet_name, filename in plot_files:
        # Figures now live under plots/png/ (SVG siblings under plots/svg/).
        img_path = plots_dir / "png" / filename
        if not img_path.is_file():
            continue
        ws = wb.create_sheet(sheet_name[:31])  # Excel sheet name limit
        try:
            img = XlImage(str(img_path))
            img.width = min(img.width, 900)
            img.height = min(img.height, 700)
            ws.add_image(img, "A1")
        except Exception as e:
            ws.cell(row=1, column=1, value=f"Could not embed {filename}: {e}")


def write_multi_project_excel_report(
    out_dir: Path,
    result: ClusteringResult,
    comparison: dict[str, Any],
    cfg: ClusteringConfig,
    *,
    project_dirs: list[Path],
    batch_correction_method: str,
) -> None:
    """Write ``clustering_report.xlsx`` for merged multi-project analysis."""
    try:
        from openpyxl import Workbook
    except ImportError:
        logger.warning(
            "openpyxl not installed — skipping Excel report. Install with: pip install openpyxl"
        )
        return

    logger.info("  Writing Excel report...")
    t0 = time.monotonic()
    styles = _excel_styles()
    if styles is None:
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    summary_rows = [
        ("Parameter", "Value"),
        ("Pipeline", "multi_project"),
        ("Output directory", str(out_dir)),
        ("Source projects", ", ".join(d.name for d in project_dirs)),
        ("Total windows", len(result.metadata)),
        ("Features per window", len(result.feature_names)),
        ("Clusters found", result.n_clusters),
        ("Clustering method", result.method),
        ("Embedding method", result.embedding_method),
        ("Batch correction", batch_correction_method),
        ("Window size (frames)", cfg.window_size),
        ("Stride (frames)", cfg.stride),
        ("k-NN neighbors", cfg.n_neighbors),
        ("Leiden resolution", cfg.leiden_resolution),
    ]
    for g, info in comparison.get("condition_groups", {}).items():
        summary_rows.append((
            f"Windows — {info.get('label', g)}",
            info.get("n_windows", 0),
        ))
    for r, (k, v) in enumerate(summary_rows, 1):
        ws.cell(row=r, column=1, value=k)
        ws.cell(row=r, column=2, value=v)
    ws.cell(row=1, column=1).font = styles["header_font"]
    ws.cell(row=1, column=2).font = styles["header_font"]
    _auto_width(ws)

    _add_cluster_sheets(wb, result, cfg, styles)

    # Condition group composition
    ws = wb.create_sheet("Condition Groups")
    cg_cols = ["Condition Group", "Label", "N Windows", "Fraction"]
    for c_idx, col_name in enumerate(cg_cols, 1):
        ws.cell(row=1, column=c_idx, value=col_name)
    _style_header(ws, len(cg_cols), styles)
    ri = 2
    for g, info in sorted(comparison.get("condition_groups", {}).items()):
        ws.cell(row=ri, column=1, value=g)
        ws.cell(row=ri, column=2, value=info.get("label", g))
        ws.cell(row=ri, column=3, value=info.get("n_windows", 0))
        ws.cell(row=ri, column=4, value=info.get("fraction", 0))
        ws[f"D{ri}"].number_format = styles["pct_fmt"]
        ri += 1
    _auto_width(ws)

    # Pairwise comparison index
    ws = wb.create_sheet("Pairwise Comparisons")
    pw_cols = [
        "Comparison Key", "Group A", "Group B", "N A", "N B",
        "Sig Features (p<0.05)", "Note",
    ]
    for c_idx, col_name in enumerate(pw_cols, 1):
        ws.cell(row=1, column=c_idx, value=col_name)
    _style_header(ws, len(pw_cols), styles)
    ri = 2
    for section_key in ("pairwise_condition_groups", "pairwise_role_conditions"):
        for key, comp in comparison.get(section_key, {}).items():
            ws.cell(row=ri, column=1, value=key)
            ws.cell(row=ri, column=2, value=comp.get("display_a", comp.get("label_a", "")))
            ws.cell(row=ri, column=3, value=comp.get("display_b", comp.get("label_b", "")))
            ws.cell(row=ri, column=4, value=comp.get("n_a", ""))
            ws.cell(row=ri, column=5, value=comp.get("n_b", ""))
            ws.cell(row=ri, column=6, value=comp.get("n_significant_features", ""))
            ws.cell(row=ri, column=7, value=comp.get("note", ""))
            ri += 1
    _auto_width(ws)

    # Housing-style enrichment for single-project backward compat in shared writer — skip

    # Top significant features from first pairwise (sample) — optional skip

    plots_dir = out_dir / "plots"
    _embed_plot_sheets(wb, plots_dir, [
        ("UMAP Clusters", "umap_by_cluster.png"),
        ("UMAP Experiment", "umap_by_experiment.png"),
        ("UMAP Condition", "umap_by_condition_group.png"),
        ("UMAP Batch", "umap_by_batch.png"),
        ("Condition Boxplots", "condition_feature_boxplots.png"),
        ("Cluster Composition", "cluster_composition_by_condition.png"),
        ("Batch Correction", "batch_correction_before_after.png"),
        ("Cluster Enrichment", "enrichment_bars.png"),
        ("Feature Enrichment", "feature_enrichment.png"),
    ])

    xlsx_path = out_dir / "clustering_report.xlsx"
    wb.save(str(xlsx_path))
    logger.info("  Wrote %s (%s)", xlsx_path.name, _fmt_size(xlsx_path.stat().st_size))
    logger.info("  Excel report done in %s", _fmt_duration(time.monotonic() - t0))


def write_single_project_excel_report(
    out_dir: Path,
    result: ClusteringResult,
    comparison: dict[str, Any],
    cfg: ClusteringConfig,
    project_name: str,
    *,
    single_animal: bool = False,
) -> None:
    """Write ``clustering_report.xlsx`` for a single-project run.

    In social mode this includes housing-enrichment sheets; in single-animal
    mode (``single_animal=True``) those are replaced by a plain cluster-sizes
    sheet and the housing-specific plots are omitted.
    """
    try:
        from openpyxl import Workbook
    except ImportError:
        logger.warning(
            "openpyxl not installed — skipping Excel report. Install with: pip install openpyxl"
        )
        return

    logger.info("  Writing Excel report...")
    t0 = time.monotonic()
    styles = _excel_styles()
    if styles is None:
        return

    sci_fmt = styles["sci_fmt"]
    float_fmt = styles["float_fmt"]
    pct_fmt = styles["pct_fmt"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    summary_rows = [
        ("Parameter", "Value"),
        ("Project", project_name),
        ("Total windows", len(result.metadata)),
        ("Features per window", len(result.feature_names)),
        ("Clusters found", result.n_clusters),
        ("Clustering method", result.method),
        ("Embedding method", result.embedding_method),
        ("Window size (frames)", cfg.window_size),
        ("Stride (frames)", cfg.stride),
        ("k-NN neighbors", cfg.n_neighbors),
        ("Leiden resolution", cfg.leiden_resolution),
        ("Normalization", cfg.normalize_method),
    ]
    overall = comparison.get("overall_enrichment", {})
    if overall:
        summary_rows.append((
            "Windows with isolated mouse",
            overall.get("total_windows_with_isolated", ""),
        ))
        summary_rows.append(("Windows group-only", overall.get("total_windows_group_only", "")))
    for r, (k, v) in enumerate(summary_rows, 1):
        ws.cell(row=r, column=1, value=k)
        ws.cell(row=r, column=2, value=v)
    ws.cell(row=1, column=1).font = styles["header_font"]
    ws.cell(row=1, column=2).font = styles["header_font"]
    _auto_width(ws)

    _add_cluster_sheets(wb, result, cfg, styles)

    per_cluster = comparison.get("per_cluster", {})
    if single_animal:
        # No enrichment in single-animal mode — just cluster sizes.
        ws = wb.create_sheet("Cluster Sizes")
        size_cols = ["Cluster", "Size", "Fraction"]
        for c_idx, col_name in enumerate(size_cols, 1):
            ws.cell(row=1, column=c_idx, value=col_name)
        _style_header(ws, len(size_cols), styles)
        n_total = len(result.cluster_labels)
        for ri, c in enumerate(range(result.n_clusters), 2):
            size = int(np.sum(result.cluster_labels == c))
            ws.cell(row=ri, column=1, value=c)
            ws.cell(row=ri, column=2, value=size)
            ws.cell(row=ri, column=3, value=(size / n_total if n_total else 0.0))
            ws[f"C{ri}"].number_format = pct_fmt
        _auto_width(ws)
    else:
        ws = wb.create_sheet("Cluster Enrichment")
        enr_cols = [
            "Cluster", "Size", "Fraction",
            "Isolated in Cluster", "Group in Cluster",
            "Isolated Outside", "Group Outside",
            "Odds Ratio", "p-value", "Significant",
        ]
        for c_idx, col_name in enumerate(enr_cols, 1):
            ws.cell(row=1, column=c_idx, value=col_name)
        _style_header(ws, len(enr_cols), styles)
        for ri, c_str in enumerate(sorted(per_cluster, key=int), 2):
            pc = per_cluster[c_str]
            enr = pc["isolated_enrichment"]
            p_val = enr["p_value"]
            ws.cell(row=ri, column=1, value=int(c_str))
            ws.cell(row=ri, column=2, value=pc["size"])
            ws.cell(row=ri, column=3, value=pc["fraction"])
            ws[f"C{ri}"].number_format = pct_fmt
            ws.cell(row=ri, column=4, value=enr["isolated_in_cluster"])
            ws.cell(row=ri, column=5, value=enr["group_in_cluster"])
            ws.cell(row=ri, column=6, value=enr["isolated_outside"])
            ws.cell(row=ri, column=7, value=enr["group_outside"])
            ws.cell(row=ri, column=8, value=enr["odds_ratio"])
            ws.cell(row=ri, column=9, value=p_val)
            ws[f"I{ri}"].number_format = sci_fmt
            ws.cell(row=ri, column=10, value=_sig_marker(p_val))
        _auto_width(ws)

    feat_enr = comparison.get("feature_enrichment", [])
    if feat_enr:
        ws = wb.create_sheet("Feature Enrichment")
        fe_cols = [
            "Feature", "Mean (Isolated)", "Mean (Group)",
            "Cohen's d", "Log2 FC", "MW p-value", "Direction", "Significant",
        ]
        for c_idx, col_name in enumerate(fe_cols, 1):
            ws.cell(row=1, column=c_idx, value=col_name)
        _style_header(ws, len(fe_cols), styles)
        for ri, fe in enumerate(feat_enr, 2):
            p_val = fe["mannwhitney_p_value"]
            ws.cell(row=ri, column=1, value=fe["feature"])
            ws.cell(row=ri, column=2, value=fe["mean_isolated"])
            ws.cell(row=ri, column=3, value=fe["mean_group"])
            ws.cell(row=ri, column=4, value=fe["cohens_d"])
            ws.cell(row=ri, column=5, value=fe["log2_fold_change"])
            ws.cell(row=ri, column=6, value=p_val)
            ws.cell(row=ri, column=7, value=fe["direction"])
            ws.cell(row=ri, column=8, value=_sig_marker(p_val))
        _auto_width(ws)

    per_feat = comparison.get("per_feature_tests", {})
    if per_feat:
        ws = wb.create_sheet("Feature Tests")
        ft_cols = ["Feature", "Comparison", "p-value", "Mean A", "Mean B", "Direction", "Significant"]
        for c_idx, col_name in enumerate(ft_cols, 1):
            ws.cell(row=1, column=c_idx, value=col_name)
        _style_header(ws, len(ft_cols), styles)
        ri = 2
        for fname in result.feature_names:
            if fname not in per_feat:
                continue
            ft = per_feat[fname]
            p_val = ft["p_value"]
            ws.cell(row=ri, column=1, value=fname)
            ws.cell(row=ri, column=2, value=ft["comparison"])
            ws.cell(row=ri, column=3, value=p_val)
            ws.cell(row=ri, column=4, value=ft["mean_a"])
            ws.cell(row=ri, column=5, value=ft["mean_b"])
            ws.cell(row=ri, column=6, value=ft["effect_direction"])
            ws.cell(row=ri, column=7, value=_sig_marker(p_val))
            ri += 1
        _auto_width(ws)

    plots_dir = out_dir / "plots"
    if single_animal:
        plot_sheets = [
            ("UMAP Clusters", "umap_by_cluster.png"),
            ("Feature Heatmap", "feature_heatmap.png"),
            ("Feature Violins", "feature_violins.png"),
        ]
    else:
        plot_sheets = [
            ("UMAP Clusters", "umap_by_cluster.png"),
            ("UMAP Housing", "umap_by_housing.png"),
            ("Cluster Composition", "cluster_composition.png"),
            ("Feature Heatmap", "feature_heatmap.png"),
            ("Cluster Enrichment Plot", "enrichment_bars.png"),
            ("Feature Enrichment Plot", "feature_enrichment.png"),
            ("Feature Violins", "feature_violins.png"),
        ]
    _embed_plot_sheets(wb, plots_dir, plot_sheets)

    xlsx_path = out_dir / "clustering_report.xlsx"
    wb.save(str(xlsx_path))
    logger.info("  Wrote %s (%s)", xlsx_path.name, _fmt_size(xlsx_path.stat().st_size))
    logger.info("  Excel report done in %s", _fmt_duration(time.monotonic() - t0))
